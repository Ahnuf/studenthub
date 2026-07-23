"""
Local, free text extraction for uploaded Note files.

No network calls, no LLM -- just parsing bytes that are already on
disk/in memory. Failures here should never block an upload; a note
with empty extracted_text just won't turn up in content search
(it's still findable by heading/course as before).
"""

import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def extract_text(file_field) -> str:
    """
    Return the best-effort extracted text for a Note's file field.

    Returns an empty string (never raises) if the extension isn't
    extractable or parsing fails -- extraction is a nice-to-have
    for search, not a hard requirement for the upload to succeed.
    """

    if not file_field:
        return ""

    extension = Path(file_field.name).suffix.lower().lstrip(".")

    try:
        if extension == "pdf":
            return _extract_pdf(file_field)

        if extension == "docx":
            return _extract_docx(file_field)

        if extension == "doc":
            # Legacy .doc (pre-2007 binary format) isn't supported by
            # python-docx. Leaving this unextracted rather than adding
            # a heavier dependency (e.g. antiword/libreoffice) for a
            # format uploaders are unlikely to use going forward.
            return ""

        if extension == "pptx":
            return _extract_pptx(file_field)

        if extension == "ppt":
            # Same reasoning as legacy .doc above.
            return ""

        return ""

    except Exception:
        logger.warning(
            "Text extraction failed for note file %s",
            file_field.name,
            exc_info=True,
        )
        return ""


def _extract_pdf(file_field) -> str:
    from pypdf import PdfReader

    file_field.seek(0)
    reader = PdfReader(file_field)

    pages = [page.extract_text() or "" for page in reader.pages]

    return "\n".join(pages).strip()


def _extract_docx(file_field) -> str:
    import docx

    file_field.seek(0)
    document = docx.Document(file_field)

    paragraphs = [p.text for p in document.paragraphs if p.text]

    return "\n".join(paragraphs).strip()


def _extract_pptx(file_field) -> str:
    from pptx import Presentation

    file_field.seek(0)
    presentation = Presentation(file_field)

    chunks = []

    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text
                if text:
                    chunks.append(text)

    return "\n".join(chunks).strip()